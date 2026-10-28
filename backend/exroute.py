from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from io import BytesIO
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font
from database import attendance_collection, subjects_collection
from depend import get_current_user

router = APIRouter(
    prefix="/excel",
    tags=["Excel"]
)

@router.get("/template")
def download_template(
    current_user=Depends(get_current_user)
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    sheet.append([
        "Subject",
        "Date",
        "Status"
    ])

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 15

    status_dropdown = DataValidation(
        type="list",
        formula1='"Present,Absent"',
        allow_blank=False
    )

    status_dropdown.error = "Select Present or Absent."
    status_dropdown.errorTitle = "Invalid Status"

    status_dropdown.prompt = "Select attendance status."
    status_dropdown.promptTitle = "Attendance"

    sheet.add_data_validation(status_dropdown)

    status_dropdown.add("C2:C1000")

    for row in range(2, 1001):
        sheet.cell(
            row=row,
            column=2
        ).number_format = "yyyy-mm-dd"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
            "attachment; filename=attendance_template.xlsx"
        }
    )


@router.get("/download")
def download_excel(
    current_user=Depends(get_current_user)
):
    user_id = str(current_user["_id"])

    records = list(
        attendance_collection.find({
            "user_id": user_id
        })
    )

    subjects = {
        str(subject["_id"]): subject["name"]
        for subject in subjects_collection.find({
            "user_id": user_id
        })
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    sheet.append([
        "Subject",
        "Date",
        "Status"
    ])

    for record in records:
        sheet.append([
            subjects.get(record["subject_id"], "Unknown"),
            record["date"],
            record["status"]
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=attendance.xlsx"
        }
    )


@router.post("/upload")
def upload_excel(
    file: UploadFile = File(...),
    current_user=Depends(get_current_user)
):
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected"
        )

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are allowed"
        )

    try:
        workbook = load_workbook(file.file)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Invalid or corrupted Excel file"
        )

    sheet = workbook.active

    headers = list(
        next(
            sheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True
            )
        )
    )

    while headers and headers[-1] is None:
        headers.pop()

    normalized_headers = [
        str(header).strip().casefold()
        if header is not None
        else ""
        for header in headers
    ]

    if normalized_headers != ["subject", "date", "status"]:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid Excel format. "
                "Required columns: "
                "Subject, Date, Status"
            )
        )

    user_id = str(current_user["_id"])

    subjects = {
        subject["name"].strip().lower(): str(subject["_id"])
        for subject in subjects_collection.find({
            "user_id": user_id
        })
    }

    imported = 0
    skipped = 0
    errors = []

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            max_col=3,
            values_only=True
        ),
        start=2
    ):

        subject_name, attendance_date, status = row

        if not subject_name:
            errors.append(
                f"Row {row_number}: Subject is missing"
            )
            skipped += 1
            continue

        if not attendance_date:
            errors.append(
                f"Row {row_number}: Date is missing"
            )
            skipped += 1
            continue

        if not status:
            errors.append(
                f"Row {row_number}: Status is missing"
            )
            skipped += 1
            continue

        subject_name = str(
            subject_name
        ).strip().lower()

        subject_id = subjects.get(
            subject_name
        )

        if not subject_id:
            errors.append(
                f"Row {row_number}: "
                f"Subject '{subject_name}' not found"
            )
            skipped += 1
            continue

        status = str(
            status
        ).strip().capitalize()

        if status not in {
            "Present",
            "Absent"
        }:
            errors.append(
                f"Row {row_number}: "
                f"Status must be Present or Absent"
            )
            skipped += 1
            continue

        attendance_date = str(
            attendance_date
        ).strip()

        existing = attendance_collection.find_one({
            "user_id": user_id,
            "subject_id": subject_id,
            "date": attendance_date
        })

        if existing:

            attendance_collection.update_one(
                {
                    "_id": existing["_id"]
                },
                {
                    "$set": {
                        "status": status
                    }
                }
            )

        else:

            attendance_collection.insert_one({
                "user_id": user_id,
                "subject_id": subject_id,
                "date": attendance_date,
                "status": status
            })

        imported += 1

    if imported:
        invalidate_user_analytics(user_id)

    return {
        "message": "Excel processed successfully",
        "imported": imported,
        "skipped": skipped,
        "errors": errors
    }